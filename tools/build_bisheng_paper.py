"""Build the five-chapter paper from reviewed content, existing photos and source excerpts."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]


def diagram(out, kind):
    im = Image.new('RGB', (1500, 900), 'white')
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 32)
    def box(x, y, w, h, label):
        draw.rounded_rectangle((x, y, x+w, y+h), radius=8, fill='#F5F7FA', outline='#344454', width=3)
        lines = label.split('\n')
        for i, line in enumerate(lines):
            b = draw.textbbox((0, 0), line, font=font)
            draw.text((x+(w-b[2])/2, y+(h-len(lines)*44)/2+i*44), line, font=font, fill='black')
    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill='#344454', width=4)
        if x2 > x1:
            points = [(x2,y2),(x2-17,y2-9),(x2-17,y2+9)]
        elif x2 < x1:
            points = [(x2,y2),(x2+17,y2-9),(x2+17,y2+9)]
        else:
            points = [(x2,y2),(x2-9,y2-17),(x2+9,y2-17)]
        draw.polygon(points, fill='#344454')
    if kind == 'architecture':
        for y, label in [(60,'扫码枪 / USB HID'),(280,'摄像头 / 图像采集'),(500,'USB 麦克风 / 音频')]:
            box(30,y,365,140,label)
            arrow(395,y+70,525,y+70)
        box(525,60,440,580,'QSM368ZP-WF\n\nQML 本机交互\nFlask API / SQLite\nRKNN CLI / NPU\n语音状态机 / 固定播报')
        for y,label in [(60,'HDMI 收银显示'),(280,'耳机 / 中文播报'),(500,'局域网 / 云演示支付')]:
            box(1100,y,365,140,label)
            arrow(965,y+70,1100,y+70)
        box(325,730,850,115,'ADB 仅调试部署　业务不依赖电脑浏览器')
    elif kind == 'scan':
        box(525,20,450,105,'HID 输入 → /api/scan')
        box(30,200,660,120,'已知条码：加购一次并刷新总价')
        box(810,200,660,120,'未知条码：显示未录入，不加购')
        arrow(640,125,360,200); arrow(860,125,1140,200)
        box(30,400,660,140,'QML → verify_scan(capture=true)\n拍照并校验已扫商品')
        box(810,400,660,140,'QML → candidates(capture=true)\n新拍照，不复用上一商品图片')
        arrow(360,320,360,400); arrow(1140,320,1140,400)
        box(30,630,660,150,'成功：Top-3 / 一致性参考\n视觉路径不再次加购')
        box(810,630,660,150,'成功：候选供人工核对\n用户点击快捷商品才加购')
        arrow(360,540,360,630); arrow(1140,540,1140,630)
        draw.text((155,835),'拍照忙 / 冷却 / 模型不可用 → 明确失败，无新结果，不用旧图顶替',font=font,fill='black')
    else:
        labels=['用户结账 → 本地未支付订单','云端可用 → 创建云端演示订单','手机打开页面 → 点击模拟支付成功','板端轮询云订单 → 本地状态更新为 paid','提交数据库事务 → QML 状态刷新 / 支付成功播报']
        for i,label in enumerate(labels):
            y=20+i*172
            box(140,y,1220,105,label)
            if i<4: arrow(750,y+105,750,y+172)
    im.save(out)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', required=True)
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    assets = out.parent / '_paper_assets'
    assets.mkdir(exist_ok=True)
    original = Document(args.source)
    photos = {}
    for name, index in [('hdmi',9),('hardware',38)]:
        rid = original.paragraphs[index]._element.xpath('.//a:blip/@r:embed')[0]
        photos[name] = original.part.related_parts[rid].blob
    for kind in ('architecture','scan','payment'):
        diagram(assets / (kind+'.png'), kind)

    doc = Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(2.0)
    sec.left_margin=Cm(2.4); sec.right_margin=Cm(2.4)
    for name,size,bold in [('Normal',11,False),('Title',20,True),('Heading 1',16,True),('Heading 2',13,True),('Heading 3',11,True),('Caption',9,False)]:
        style=doc.styles[name]
        style.font.name='SimSun' if name=='Normal' else 'Microsoft YaHei'
        style.font.size=Pt(size); style.font.bold=bold; style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),style.font.name)
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.35
    doc.styles['Normal'].paragraph_format.first_line_indent=Cm(.75)
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
    def para(text, style=None):
        return doc.add_paragraph(text,style)
    def head(text, level=2):
        return doc.add_heading(text,level)
    def page():
        doc.add_page_break()
    def picture(data, caption, width=15):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(io.BytesIO(data) if isinstance(data,bytes) else str(data),width=Cm(width))
        p=para(caption,'Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    def table(headers, rows, widths):
        t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
        for i,w in enumerate(widths): t.columns[i].width=Cm(w)
        for data,row in [(headers,t.rows[0])]+[(r,t.add_row()) for r in rows]:
            for i,value in enumerate(data):
                cell=row.cells[i]; cell.width=Cm(widths[i]); cell.vertical_alignment=1
                tcPr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
                tcPr.append(borders)
                margin=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margin.append(el)
                tcPr.append(margin)
                p=cell.paragraphs[0]; p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.15
                r=p.add_run(str(value)); r.font.size=Pt(9.5)
                if data is headers:
                    r.bold=True; sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF1'); tcPr.append(sh)
            trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        para('')
    def code(text):
        p=doc.add_paragraph()
        p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.line_spacing=1.1
        p.paragraph_format.keep_together=True
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F3F4F5'); p._p.get_or_add_pPr().append(shade)
        r=p.add_run(textwrap.dedent(text).strip()); r.font.name='Consolas'; r.font.size=Pt(8.5)

    # Page 1
    p=para('毕昇杯初赛设计论文'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=para('基于嵌入式系统的条码及端侧视觉校验的自助收银终端','Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for txt in ['学校  北京林业大学','院系或专业  电子信息科学与技术','参赛成员  张正  吴永嘉  管湘雪  龚思颖','指导教师  张立','2026 年 9 月']:
        p=para(txt); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
    head('摘要')
    para('面向校园便利点、宿舍零售点和办公室自助结算场景，设计一套基于 QSM368ZP-WF 嵌入式 Linux 平台的自助收银终端。系统通过 HDMI 输出 Qt/QML 本机界面，以 Flask API 和 SQLite 管理商品、购物车及订单。条码承担稳定计费，摄像头图像经 RKNN C runtime 桥接程序在端侧 NPU 上进行 10 类商品分类，给出 Top-3 候选与一致性参考。视觉与加购分离，避免一次扫码后又因推理结果重复计费；未知条码触发新拍照并供用户核对。语音采用唤醒词与编号短语，固定中文 TTS 音频提供操作反馈。支付部分使用云端模拟页面和板端轮询，验证订单状态同步，不涉及真实资金结算。')
    para('历史固定测试集共 18 张图像，Top-1 为 17/18，Top-3 为 18/18；该结果属于低样本验证，不能外推为现场泛化准确率。本次修订以隔离测试验证未知条码的新图获取与失败分支，保留原有板端启动和恢复方式。作品的重点是可靠计费与辅助感知的职责分离及端侧完整业务链路。')
    para('关键词：嵌入式 Linux；条码收银；RKNN；端侧视觉校验；人机交互')

    # Page 2
    page(); head('一、设计题目与研究背景',1); head('1.1 设计题目')
    para('设计题目为“基于嵌入式系统的条码及端侧视觉校验的自助收银终端”。作品形态为桌面式自助结算原型，不是自动出货售货机；不包含货道、柜门锁、称重或库存传感器。目标对象为商品种类有限、用户自行放置商品并主动扫码的小型零售点。')
    head('1.2 背景及意义')
    para('自助结算需要在交互便利性、计价确定性和设备成本之间取得平衡。仅依靠商品外观决定收费，容易将分类错误直接传递到金额；仅使用条码又缺少对当前商品外观的辅助提示。本项目把二者分离：条码命中商品表后记账，视觉只给出参考，不更改条码价格、不自动追加商品。该设计便于解释计费来源，也允许摄像头或模型暂时不可用时继续显示购物车。')
    para('本机 QML 界面、端侧推理和 SQLite 使主要交互留在开发板上。云端仅参与模拟支付状态中转；若现场借用电脑共享网络，电脑承担网络出口而非收银运算。系统不把本地调试地址当作手机可访问支付链接。')
    head('1.3 研究现状')
    para('从技术路线看，可比较条码输入、图像分类和融合辅助三种方案。图像分类为整张图片输出预定义类别，不提供多目标位置与数量，因此当前 YOLOv8n-cls 不适合一次识别整篮商品 [1]。RKNN 工具链提供模型转换与设备运行接口，为将分类模型部署到 Rockchip NPU 提供条件 [2]。本项目没有提出新型神经网络，创新主要在系统职责划分与可恢复的工程流程。')
    table(['路线','优点','边界'],[
        ['条码主通道','商品编号可直接关联价格','依赖标签可读和条码库正确'],
        ['纯视觉计费','无需主动扫描标签','误分类直接影响收费；需更多场景验证'],
        ['本项目融合方式','条码计费，视觉供复核','不构成防盗系统；不保证发现所有错贴码'],
    ],[3,6,7.2])
    para('Qt Quick 提供 QML 声明式界面能力 [3]，Flask 作为轻量 WSGI 框架连接业务接口 [4]，SQLite 支持事务式本地存储 [5]。这些成熟组件减少了自行实现底层框架的工作，但接口幂等、物理断电恢复和多终端并发仍需单独验证。')

    # Page 3
    page(); head('二、硬件介绍',1); head('2.1 系统总体设计')
    para('以 QSM368ZP-WF 开发板为中心，使用 RK3568 类 aarch64 Linux 环境、Weston/Wayland 显示链路和已部署在 /userdata/qt5 的 Qt 运行时。显示器、扫码枪、摄像头和 USB 音频构成输入输出层；以太网连接外部网络。ADB 仅用于诊断、部署和取证，不承担面向顾客的操作。')
    picture(assets/'architecture.png','图 1 系统总体结构与端云职责',15.6)
    para('设计采用现有开发板与标准外设，不自制 PCB，不增设机械执行器。硬件介绍重点是接口角色、供电约束、模块协作和实测连接，不能把接口集成写成自行设计了底层 SoC 电路。')
    table(['连接','用途','工程要求'],[
        ['HDMI → 显示器','QML 本机输出','确认 connected、输出位置及界面肉眼可见'],
        ['USB Host → 外设','HID 扫码、摄像头、音频','固定已验证端口；与 OTG/ADB 区分'],
        ['Ethernet → 网络','模拟云订单访问','检查链路、地址、路由及云 /health'],
    ],[4.2,4.5,7.5])

    # Page 4
    page(); head('2.2 系统主要功能模块设计')
    table(['模块','组成及接口','设计边界'],[
        ['控制与存储','开发板、Linux、SQLite','应用层更新，不修改 boot/rootfs'],
        ['扫码输入','USB HID 扫码枪','扫描成功蜂鸣不代表 QML 已收到输入'],
        ['图像采集','USB 摄像头、GStreamer','单商品放置；失败/冷却必须可见'],
        ['显示交互','HDMI 显示器、Qt/QML','随屏幕分辨率检验布局和支付弹窗'],
        ['语音输入输出','HyperX Cloud III、ALSA','固定设备名；缺失时不阻断主业务'],
        ['网络与调试','以太网、ADB 数据线','网络中断不应被误报为支付成功'],
    ],[3.3,5.8,7.1])
    picture(photos['hardware'],'图 2 原项目实物连接照片（历史展示，不作为本次接线验收）',10.5)
    para('原项目曾出现扫码枪供电及蜂鸣正常、但所接端口未把 HID 事件送入界面的情况。现场应固定扫码、摄像头、音频和调试端口并贴标签。供电采用与开发板规格匹配的适配器；电压、电流和外设总功耗需实测后填写，本文不编造数值。')
    para('成本统计范围包括开发板 1 套、显示器 1 台、扫码枪 1 把、摄像头 1 个、USB 音频 1 套及电源/线材。采购单价、总成本和持续运行功耗尚未形成可核验记录，因此目前只能说明组件组成，不能声称已证明成本优势。')

    # Page 5
    page(); head('三、软件设计',1); head('3.1 程序流程图')
    para('QML 通过 retail_api.js 调用 JSON API，Flask 接口管理业务状态，SQLite 保存商品与订单。图 3 区分已知条码与未知条码两条路径：/api/scan 只负责校验和加购；自动拍照由 QML 收到结果后发起视觉请求，不是扫码 API 自身的副作用。')
    picture(assets/'scan.png','图 3 扫码计费及未知条码新拍照流程',16)
    para('候选流程增加 capture=true：无论之前是否显示照片，都调用现有拍照锁、冷却和超时机制获取新图。失败时返回空候选和错误原因，不使用上一商品图片顶替；模型失效时也不把模拟候选当作真实推理。新图可用于预览，但其分类失败不会改变购物车。')
    para('QML 展示候选后，由用户核对商品并点击相应快捷商品加购。该方式不是自动绑定未知条码，也不为未录入 SKU 创建价格。已有 /api/vision/confirm 接口可用于显式人工确认，但本次不增加新的候选弹窗。')

    # Page 6
    page(); head('3.1 程序流程图（续）')
    picture(assets/'payment.png','图 4 仿真支付与板端轮询同步流程',13.3)
    para('手机页面的“模拟支付成功”只改变演示订单状态。板端通过 cloud_fetch_status 获取云订单，sync_order_from_cloud 更新本地记录后提交并关闭数据库连接，再触发 payment_success 播报。这是轮询同步，不是银行或支付平台的签名回调，也不代表发生真实资金交易。')
    para('支付 URL 按 cloud、lan、local_debug 区分可访问范围。局域网模式要求手机与板端可互访；127.0.0.1 只指向访问设备自身，不能当作手机付款入口。云不可达时保留未支付状态，不自动视为已付款。')
    para('语音流程收敛为“唤醒词→编号短语→意图检查→执行或取消”。菜单覆盖总价、删除、清空、结账与拍照；危险操作的确认依具体入口实现。中文播报采用预生成的神经 TTS PCM WAV，不是现场真人录音，也不要求操作时联网生成语音。')

    # Page 7
    page(); head('3.2 核心代码片段')
    para('片段均来自本次本地源代码；仅节选相关逻辑，省略部分未改变执行顺序。完整返回字段、错误分支和设备调用见 app.py 与 qt_kiosk/。')
    head('条码计费及自动视觉调用',3)
    code('''
        product = find_product_by_barcode(raw_code)
        if not product:
            product = get_product(raw_code.upper())
        # 未命中：返回 unknown_barcode，不执行以下加购
        # 已命中：
        add_product(product["product_id"])
        log_recognition("barcode_scanner", True, raw_code=raw_code,
                        product=product, confidence=1.0, latency_ms=latency)
        play_audio("scan_success")
    ''')
    para('上段为 /api/scan 节选。QML 已知条码成功分支再调用 visionVerifyScan(code, "", true, callback)，未知条码分支改为 visionCandidates("", true, callback)。不向新拍照请求传旧的 capturePreviewPath。')
    head('新图采集失败即停止候选流程',3)
    code('''
        fresh_capture = data.get("capture") is True
        cap = {}
        if fresh_capture:
            # 先更新候选状态为 capturing
            cap = api_capture_v265().get_json(silent=True) or {}
            if not cap.get("ok") or not cap.get("image_path"):
                # 构造 ok=false、空 image_path、空 candidates 的 payload
                set_latest_vision(payload)
                return jsonify(payload)
            image_path = cap["image_path"]
        else:
            image_path = data.get("image_path") or get_state("latest_capture", "")
    ''')
    para('capture 必须为 JSON 布尔值 true 才启动新流程。未传 capture 的旧调用仍可指定 image_path 或读取历史照片，保持既有诊断工具兼容。只有新流程禁止失败后降级到旧图。')
    head('候选不是计费指令',3)
    code('''
        real_prediction = bool(predicted.get("ok") and
                               predicted.get("model_available") and candidates)
        if fresh_capture and not real_prediction:
            candidates = []
        # payload: auto_add_cart=False；新流程 ok=real_prediction
    ''')

    # Page 8
    page(); head('3.2 核心代码片段（续）')
    head('订单状态与播报分离',3)
    lines=(ROOT/'app.py').read_text(encoding='utf-8').splitlines()
    begin=next(i for i,s in enumerate(lines) if s.startswith('def sync_order_from_cloud('))
    end=next(i for i in range(begin,len(lines)) if lines[i].startswith('def init_db('))
    tail=lines[begin:end]
    idx=next(i for i,s in enumerate(tail) if 'conn.commit()' in s)
    code('\n'.join(tail[idx:idx+4]))
    para('这是 sync_order_from_cloud 中的真实连续片段。先释放数据库事务再播报，避免把音频调用混入订单写事务。它能减少事务与设备调用耦合，但不等于已经完成真实商户回调验签、退款或断电恢复验证。')
    head('端侧模型与故障边界',3)
    para('视觉路径先将图像预处理为 224×224 RGB 输入，再调用 C++ CLI 加载 RKNN 模型，通过 librknnrt.so 执行并返回 JSON。Flask 解析类别映射与 Top-3；RKNN 路径异常时保留 ONNX Runtime 兼容路径。模型分数在界面标为“参考分”，不当作经过校准的计费正确概率。')
    table(['约束','实现策略','不能据此宣称'],[
        ['重复收费','视觉接口不自动加购','所有 HTTP 重试、双击都已幂等'],
        ['误识别','参考结果与人工核对分离','能识别所有错贴码或防止盗损'],
        ['照片过期','未知条码强制新拍照，失败不读旧图','所有并发跨商品结果均已逐件关联'],
        ['设备不可用','拍照/音频错误可见，保留扫码计费','USB 热插拔或断电绝不影响系统'],
    ],[3.1,6.1,7])
    para('初赛阶段保留单体后端，不重构微服务，不增加库存传感器、连续视觉自动加购或真实支付 SDK。后续优先补充请求标识、候选与扫描事件关联以及提交幂等测试，再扩展场景。')

    # Page 9
    page(); head('四、测试与验证',1); head('4.1 数据集与指标口径')
    para('历史训练数据共 10 类、117 张图，按固定随机种子 42 近似 70/15/15 划分，实际为训练 81 张、验证 18 张、测试 18 张。模型为 YOLOv8n-cls，历史训练配置为 50 epochs、输入尺寸 224。此处复核既有预测 CSV 与报告，不声称本次重新训练或重新执行 NPU 评测。')
    rows=[]
    with (ROOT/'dataset_raw_v260/manifest.csv').open(encoding='utf-8-sig',newline='') as f: manifest=list(csv.DictReader(f))
    summary=json.loads((ROOT/'MODEL_ARTIFACTS_V260/evaluation_summary.json').read_text(encoding='utf-8'))
    for m,stats in zip(manifest,summary['per_class']):
        n=stats['total']; correct=round(stats['top1_acc']*n)
        rows.append([m['product_id'],m['product_name'],m['image_count'],str(n),f'{correct}/{n}',f'{n}/{n}'])
    table(['SKU','商品','总图数','测试图','Top-1','Top-3'],rows,[2,5.4,2,2,2.4,2.4])
    para('以测试集类别标签为真值，Top-1=17/18=94.44%，Top-3=18/18=100%。唯一 Top-1 失败类是 SKU005，该类测试样本仅 2 张。Top-3 表示真值落入前三候选，不代表第一候选总是正确。每类测试只有 1～3 张，指标对划分与拍摄条件敏感。')
    para('目前未提供按独立拍摄批次或商品实例分组的外部测试集；相似背景和连续拍摄可能使结果偏乐观。下一步应增加不同日期、光照、角度和背景，并按采集批次隔离测试集，单独报告误匹配和低置信度比例。')
    head('4.2 转换一致性与时延')
    para('历史 RKNN Toolkit2 模拟器与 ONNX 对 117 张图进行预测一致性比较：FP 模型 Top-1 一致 117/117，INT8 一致 116/117。该比较不是识别准确率，包含全数据集，也不是 117 次板端实测。历史板端单图示例记录 CLI 约 15 ms、Flask 路径约 24 ms；不含完整扫码拍照链路，不当作平均值、P95 或帧率。')

    # Page 10. Page 9 can be full; let Word flow naturally here so a
    # trailing page-break paragraph cannot be pushed onto a blank page.
    head('4.3 系统历史验证与实物证据')
    table(['项目','历史记录','口径与限制'],[
        ['扫码字段与异常输入','39/39，2026-06','scan_api_matrix；接口断言数，不是真实扫码次数'],
        ['10-SKU 条码目录','43/43，2026-06','scan_v260_10sku；名称、条码、金额一致性'],
        ['QML 业务 API','45/45，V2.6.11-G','API-backed 路径，不等于触屏与 HID 物理验收'],
        ['音频事件','24/24，V2.6.11-G','接口/播放状态断言；听感需人工确认'],
        ['演示支付','历史云 paid 回写；后续存在网络降级','新云节点与手机完整回写仍需现场复测'],
        ['恢复与显示','已有 start/restore 历史记录','本次未操作板端，不把历史 PASS 写成当前 PASS'],
    ],[3.4,5.0,7.8])
    picture(photos['hdmi'],'图 5 原项目 HDMI 收银界面照片（历史实物证据）',12.5)
    para('历史数据来源为 FINAL_TEST_RESULTS.md、V2611G_FINAL_TEST_RESULTS.md 和对应专题报告。原 Word 的照片可说明设备形态和界面存在，但不能证明当前补丁已部署。最终现场复测须记录版本、日期、USB 端口、网络拓扑及订单编号，避免跨版本混用。')

    # Page 11
    page(); head('4.4 本次补丁的隔离验证')
    para('2026 年 9 月 22 日在 Windows 主机复制 app.py 到临时目录，以独立 SQLite 数据库和 Flask test_client 执行 12 项 unittest；摄像头、推理及播放使用测试替身，云端创建订单被禁用。这样验证应用分支而不触碰运行中板端数据。结果为 12/12 通过。')
    table(['验证组','检查内容','结果'],[
        ['新图与失败分支','强制新图；忙/冷却/超时/失败不推理旧图；缺路径拒绝','通过'],
        ['真实候选约束','模型不可用、模拟候选、空输出不显示为成功','通过'],
        ['兼容与计费','旧调用兼容；未知不加购；已知校验不重复加购','通过'],
        ['主业务回归','加购/金额/删除/清空；空结账保护；本地订单及 QR','通过'],
        ['语音与 UI 逻辑','文本查询/纠错；JS 新参数、反馈、预览、焦点恢复','通过'],
    ],[3.1,10.8,2.3])
    code('''
        python tools/test_unknown_barcode_capture.py
        node tools/test_unknown_barcode_qml.js
    ''')
    para('JS 测试在 Node vm 中执行 Main.qml 的实际 submitScan 函数，并使用 API/UI 替身；它不渲染 Qt，不证明相机驱动、NPU 实际执行或 HDMI 显示。因此本次结论为本地逻辑验证通过，硬件现场验收待执行。')
    head('4.5 现场验收方案')
    table(['场景','通过标准'],[
        ['已知 SKU001 / SKU002','各扫码一次，各加购一次；金额等于商品表价格之和'],
        ['未知条码与换商品','先拍 A，再放 B 扫未知码；必须显示 B 新图与候选，不加购'],
        ['候选失败与恢复','冷却/拔摄像头/模型失效有提示，不把 A 旧图结果冒充 B'],
        ['人工加购与付款','核对快捷商品后加入一次；模拟付款后本地 paid，不重复播报'],
        ['恢复与证据','start/restore 回环；保存 HDMI、接线、手机页面与日志'],
    ],[4.5,11.7])

    # Page 12
    page(); head('五、结论',1)
    para('作品已形成以嵌入式 Linux 开发板为中心的自助结算原型：扫码承担确定性计费，QML 提供本机收银界面，RKNN/NPU 分类提供视觉参考，语音菜单与中文播报降低操作门槛，仿真云支付验证订单状态同步。该方案的贡献在于把辅助感知与收费动作解耦，并对失效情况保留可见反馈和恢复路径，而非宣称提出新的识别算法或完成商业支付系统。')
    para('本次改进保留已验证主流程，只补齐未知条码的新图获取和失败反馈，避免陈旧图像造成候选误导。12 项隔离测试及 JS 分支测试通过；板端新版本摄像头、NPU、HDMI 和手机支付联测仍需现场完成。小样本分类指标与历史转模型一致性测试均按各自口径披露，不用于保证真实环境的安全性或识别准确率。')
    para('后续优先完成固定接线、稳定网络、单商品放置区域、带版本的连续演示视频和 BOM 记录；再扩展跨拍摄批次数据、端到端 P50/P95 时延、长期运行与断电恢复实验。真实商户支付、库存传感器、多商品检测和一体化外壳属于后续产品化范围，不作为本次初赛已实现成果。')
    head('参考资料')
    refs=[
        '[1] Ultralytics. Image Classification. https://docs.ultralytics.com/tasks/classify/ （2026-09-22 查阅；用于分类任务定义，项目仍用历史 YOLOv8n-cls）。',
        '[2] Rockchip. RKNN-Toolkit2. https://github.com/airockchip/rknn-toolkit2 （2026-09-22 查阅）。',
        '[3] The Qt Company. Qt Quick 5.15. https://doc.qt.io/archives/qt-5.15/qtquick-index.html （2026-09-22 查阅）。',
        '[4] Pallets. Flask Documentation. https://flask.palletsprojects.com/en/stable/ （2026-09-22 查阅）。',
        '[5] SQLite. SQLite Is Transactional. https://www.sqlite.org/transactional.html （2026-09-22 查阅）。',
        '[6] 项目内部证据：CLASSIFICATION_DATASET_BUILD_REPORT.md；MODEL_ARTIFACTS_V260/evaluation_summary.json、predictions.csv；V263_RKNN_ONNX_CONSISTENCY_REPORT.md；V263_RKNN_NPU_PERFORMANCE_REPORT.md；V2611G_FINAL_TEST_RESULTS.md。',
        '[7] 项目源码与复现：app.py；qt_kiosk/；tools/test_unknown_barcode_capture.py；tools/test_unknown_barcode_qml.js。https://github.com/zhangzheng-debug/Team11782-Smartvendingterminal 。',
    ]
    for ref in refs:
        p=para(ref); p.paragraph_format.first_line_indent=Cm(0)
        for r in p.runs:r.font.size=Pt(9)
    doc.core_properties.title='基于嵌入式系统的条码及端侧视觉校验的自助收银终端'
    doc.core_properties.author='张正 吴永嘉 管湘雪 龚思颖'
    doc.core_properties.subject='毕昇杯初赛 五章结构设计论文'
    doc.save(out)
    print(out)
    print('sha256',hashlib.sha256(out.read_bytes()).hexdigest())


if __name__ == '__main__':
    main()
